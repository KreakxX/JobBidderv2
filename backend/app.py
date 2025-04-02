from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import uvicorn
import undetected_chromedriver as uc
import time
from selenium.webdriver.common.by import By
from job import job
from typing import List
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.action_chains import ActionChains
from docx import Document


app = FastAPI()

cache_LinkedIn = []
cache_Indeed = []

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class Cached: 
    def __init__(self, arr: List[job], title: str, ort: str):
        self.arr = arr
        self.title = title
        self.ort = ort
        self.created = int(time.time())

class jobRequest(BaseModel):
   title: str
   ort: str

class Job(BaseModel):
   title: str
   company: str
   url: str
   payment: str
   JobType: str
   description: str

class ResumeRequest(BaseModel):
    vorname: str
    nachname: str
    Birtday: str
    BirthPlace: str
    Nationality: str
    marital: str
    educationDate1: str
    educationDate2: str
    educationDate3: str
    trainingDate1: str
    trainingDate2: str
    trainingDate3: str
    education1: str
    education2: str
    education3: str
    training1: str
    training2: str
    training3: str
    language1: str
    language2: str
    language3: str
    programming1: str
    programming2: str
    programming3: str
    programming4: str
    programming5: str
    programming6: str
    programming7: str
    programming8: str
    programming9: str
    datum: str
    ort: str





def generatePerfectResume(request: ResumeRequest):
    template_path = 'backend/templates/LebensLaufVorlage.docx'
    document = Document(template_path)
    
    placeholders = {
        '[[Vorname]]': request.vorname,
        '[[Nachname]]': request.nachname,
        '[[BirthDate]]': request.Birtday,
        '[[Birthplace]]': request.BirthPlace,
        '[[Nationality]]': request.Nationality,
        '[[Marital Status]]': request.marital,
        '[[Date 1]]': request.educationDate1,
        '[[Date 2]]': request.educationDate2,
        '[[Date 3]]': request.educationDate3,
        '[[Date 4]]': request.trainingDate1,
        '[[Date 5]]': request.trainingDate2,
        '[[ Date 6]]': request.trainingDate3,
        '[[Education 1]]': request.education1,
        '[[Education 2]]': request.education2,
        '[[Education 3]]': request.education3,
        '[[Training 1]]': request.training1,
        '[[Training 2]]': request.training2,
        '[[Training 3]]': request.training3,
        '[[Language 1]]': request.language1,
        '[[Language 2]]': request.language2,
        '[[Language 3]]': request.language3,
        '[[Programming 1]]': request.programming1,
        '[[Programming 2]]': request.programming2,
        '[[Programming 3]]': request.programming3,
        '[[Programming 4]]': request.programming4,
        '[[Programming 5]]': request.programming5,
        '[[Programming 6]]': request.programming6,
        '[[Programming 7]]': request.programming7,
        '[[Programming 8]]': request.programming8,
        '[[Programming 9]]': request.programming9,
        '[[Datum]]': request.datum,
        '[[Ort]]': request.ort,
    }

    for p in document.paragraphs:
        for placeholder, value in placeholders.items():
            if placeholder in p.text:
             if(value == "None"):
                 p.text = p.text.replace(placeholder, "")
             else:
                p.text = p.text.replace(placeholder, value)

        
    
    document.save(f"GenerateLebenslauf_{request.vorname}.docx")


def downloadAllJobLinks(jobs: List[Job]):
    urls = []
    for job in jobs:
       url = job.url
       urls.append(url)
   
    with open('job_links.txt', 'w') as file:
         for url in urls:
               file.write(url + '\n')



def scrappeFromLinkedIn(title: str, ort: str):
    print("Geht")
    jobs = []
    driver = uc.Chrome(version_main=134, use_subprocess=False)
    driver.get(f"https://www.linkedin.com/jobs/search?keywords={title}&location={ort}&geoId=&trk=public_jobs_jobs-search-bar_search-submit&position=1&pageNum=0")
    print(driver.current_url)
    time.sleep(5)
    
    driver.execute_script("document.elementFromPoint(100, 100).click();")
    try:
      button = driver.find_element(By.CSS_SELECTOR, "button.modal__dismiss[aria-label='Verwerfen']")
      button.click()
    except:
        print("Error")
    time.sleep(2)
    try:
      button2 = driver.find_element(By.CSS_SELECTOR, 'button[data-tracking-control-name="public_jobs_f_WT"]')
      button2.click()
      time.sleep(1)
      button5 = driver.find_element(By.XPATH, "//button[@class='filter__submit-button']")
      button5.click()
      time.sleep(4)
    except:
        print("Error 2")

    try:
      checkbox = driver.find_element(By.ID,"f_WT-2")
      checkbox.click()
      time.sleep(1)
      submit_button = driver.find_element(By.CSS_SELECTOR, "button.filter__submit-button")
      submit_button.click()
      time.sleep(3)
    except:
        print("Error 3")
    time.sleep(2)
    ablehnen_button = driver.find_element(By.XPATH, "//button[contains(@class, 'artdeco-button--primary') and contains(text(), 'Ablehnen')]")
    ablehnen_button.click()
    time.sleep(1)
    last_height = driver.execute_script("return document.body.scrollHeight")
    while True:
        driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
        time.sleep(1)
        new_height = driver.execute_script("return document.body.scrollHeight")
        if new_height == last_height:
            break
        last_height = new_height
        driver.execute_script("window.scrollTo(0, 0);")
    
    elements = driver.find_elements(By.CSS_SELECTOR,"div.base-card.relative.w-full")
    for element in elements:
        if(len(jobs) > 20):
            y = Cached(jobs, title, ort)
            cache_LinkedIn.append(y)
            driver.quit() 
            return jobs;
        try:
            JobTitle = element.find_element(By.CSS_SELECTOR,".base-search-card__title").text
        except:
            title  = "Unknown"

        try:
            company = element.find_element(By.CSS_SELECTOR,".hidden-nested-link").text
        except: 
            company = "Unknown"

        try:
            url = element.find_element(By.CSS_SELECTOR, "a[data-tracking-control-name='public_jobs_jserp-result_search-card']").get_attribute("href")
        except:
            url = "Unknown"
        
        time.sleep(1)
        element.click()
        time.sleep(2) 
        try:
            description_element = driver.find_element(By.CSS_SELECTOR, ".show-more-less-html__markup.relative.overflow-hidden")
            description = description_element.text
        except:
            print("Could not find description element")

        x = job(JobTitle, company, url, "Unknown", "Unknown", description)
        jobs.append(x)
    time.sleep(2)
    y = Cached(jobs, title, ort)
    cache_LinkedIn.append(y)
    driver.quit() 
    return jobs

def scrappeFromIndeed(title: str, ort: str):
    jobs = []
    driver = uc.Chrome(version_main=134, use_subprocess=False)
    driver.get(f"https://www.indeed.com/jobs?q={title}&l=remote&fromage=1&sc=0kf%3Aattr%28DSQF7%29%3B&from=searchOnDesktopSerp")
    print(driver.current_url)
    time.sleep(10)  
    print(driver.title)
    notatend = True
    while notatend :
      elements = driver.find_elements(By.CSS_SELECTOR,".cardOutline.tapItem.dd-privacy-allow.result")
    
      for element in elements:
            Jobtitle = element.find_element(By.CSS_SELECTOR,"h2.jobTitle span").text
            company = element.find_element(By.CSS_SELECTOR,"span[data-testid='company-name']").text
            url = element.find_element(By.CSS_SELECTOR,"a[data-jk]").get_attribute("href")
            element.click()
            time.sleep(2.8)
            try:
                newElement = driver.find_element(By.CSS_SELECTOR, ".jobsearch-JobComponent.css-1kw92ky.eu4oa1w0")
            except Exception as e:
                print(f"An error occurred: {e}")
            try:
                payment = newElement.find_element(By.CSS_SELECTOR, "li[data-testid='list-item']").text
            except Exception as e:
                payment = "Unknown"
            if "$" not in payment:
                payment = "Unknown"
            try:
                JobType = newElement.find_element(By.CSS_SELECTOR,"[data-testid='Full-time-tile']").text
            except Exception as e:
                JobType = "Unknown"

            description = newElement.find_element(By.CSS_SELECTOR,".jobsearch-JobComponent-description.css-wppltw.eu4oa1w0").text
                
  
            print(payment)
            x = job(Jobtitle, company, url, payment, JobType, description)
            jobs.append(x)
      try:
         nextpage = driver.find_element(By.CSS_SELECTOR, 'a[data-testid="pagination-page-next"]')
         if nextpage:
            nextpage.click()
            time.sleep(3)
      except Exception as e:
         notatend = False   
         

    y = Cached(jobs, title, ort)
    cache_Indeed.append(y)
    time.sleep(2)  
    driver.quit()
    return jobs
 
@app.post("/scrapeJobsFromLinkedIn")
def ScrapeJobsFromLinkedIn(request: jobRequest):
    for x in cache_LinkedIn[:]:
        if (x.created + 8200) <  int(time.time()):
            cache_LinkedIn.remove(x)
    for x in cache_LinkedIn[:]:
        print(x.title)
        print(x.ort)
        print(x.arr)
        if x.title.lower() == request.title.lower() and x.ort.lower() == request.ort.lower():
            return x.arr
    return scrappeFromLinkedIn(request.title, request.ort)

@app.post("/ScrappeJobsFromIndeed")
def ScrappeJobsFromIndeed(request: jobRequest):
    for x in cache_Indeed[:]:
        if (x.created + 8200) <  int(time.time()):
            cache_Indeed.remove(x)
    for x in cache_Indeed[:]:
        print(x.title)
        print(x.ort)
        print(x.arr)
        if x.title.lower() == request.title.lower() and x.ort.lower() == request.ort.lower():
            return x.arr
    return scrappeFromIndeed(request.title, request.ort)


@app.post("/download/all/links")
def downloadAllLinks(jobs: List[Job]):
   downloadAllJobLinks(jobs);

@app.post("/generate/perfect/Resume")
def GeneratePerfectResume(request: ResumeRequest):
    generatePerfectResume(request)



if __name__ == "__main__":
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)


